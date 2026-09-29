# -*- coding: utf-8 -*-
"""Zero-residue scan over the BUILT docx files (v2.2/v3.0):
  1) stale tokens that must not appear anywhere;
  2) spot-check key fixed values are present.
"""
import os
import sys

sys.path.insert(0, os.path.join('..', '.pkgs', 'docx_pkg'))
sys.path.insert(0, os.path.join('..', '.pkgs', 'lxml_pkg'))
from docx import Document

STALE = ['70.5', '基线 80.4', '"80.4"', '4.0ms', '与 1.50',
         '0.063', '0.074', '中位数进一步压至', '四成种子', '正在整理', '[待测', '[待补',
         '一百至三百条', '增益随数据量增加而收窄', '每类三百条", "门控',
         '每类三百条", "SMOTE', '每类三百条", "真实', '降至 0.21',
         '压至 0.21', '0.21（']
# NOTE: bare '0.21' is NOT stale any more: E1 per-seed drift list contains
# -0.219 legitimately.  Old mean-quotes were '降至 0.21'/'压至 0.21'/'0.21（'.
# careful: 0.21 / 1.50 could legitimately appear in per-seed values?
# appendix A/B contain e.g. 0.71...; "1.50" was only the gate range;
# "0.21" was the double-rounded flood mean. Check context later if hit.
WANT_V22 = ['每类二百条 · se 71K', '每类一百条 · light 49K', '3.9ms',
            '28.5', '1.67', '干净增强', 'seed0', '0.20（0.01 至 0.46）',
            '80.3', '70.4']
WANT_V3 = ['每类二百条 · se 71K', '每类一百条 · se 71K',
           '每类一百条 · light 49K', '4/5 种子为正', '57.67±5.01',
           '53.18±1.86', '0.073', 'wide 主干、每类三百条门控',
           '同深度重建对照组', 'Wilcoxon p=0.004',
           '0.0111（5/5 种子为正）', '图 1', '图 2', '图 3', '参考文献',
           '\\mathrm{Acc}',
           # 6.7 深化部署评估（edge_deploy JSON）
           '1.534 M', '81.89±0.88', '+0.28±0.16', '76.2 KB', '2.51 ms',
           '96.6 ms', '149.6 ms', '30.7 MHz', '分析模型',
           # 7 跨域实例表（cicids_xgb_wgan_rerun JSON）
           '0.631→0.023', '93.5%', '0.296%', 'Web_Attack_Brute_Force']


def all_text(path):
    doc = Document(path)
    texts = [p.text for p in doc.paragraphs]
    for t in doc.tables:
        for row in t.rows:
            for cell in row.cells:
                texts.append(cell.text)
    return '\n'.join(texts)


def count_images(path):
    doc = Document(path)
    return sum(1 for rel in doc.part.rels.values()
               if 'image' in rel.reltype)


for name, path, stale, want in [
        ('v2.2', 'docs/论文_医学IoT增强投毒与防御_v2.2.docx', STALE, WANT_V22),
        ('v3.0', 'docs/论文_医学IoT增强投毒与防御_v3.0.docx', STALE, WANT_V3)]:
    txt = all_text(path)
    bad = [s for s in stale if s in txt]
    miss = [s for s in want if s not in txt]
    print('== %s ==' % name)
    print('  stale tokens found:', bad if bad else 'NONE')
    print('  wanted values missing:', miss if miss else 'NONE')
    nimg = count_images(path)
    print('  embedded images: %d (v3.0 expects 3)' % nimg)
